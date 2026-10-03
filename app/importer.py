from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from .database import get_db


class ImportValidationError(ValueError):
    pass


SHEETS = {
    "PurchaseOrders": [
        "po_number", "supplier_code", "supplier_name", "order_date",
        "expected_date", "product_sku", "product_name", "quantity", "unit_price",
    ],
    "GoodsReceipts": [
        "receipt_number", "po_number", "receipt_date", "product_sku", "quantity",
    ],
    "Invoices": [
        "invoice_number", "po_number", "supplier_code", "invoice_date",
        "product_sku", "quantity", "unit_price",
    ],
}


def _text(value):
    return str(value).strip() if value is not None else ""


def _date(value, label):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _text(value)
    try:
        return datetime.strptime(raw, "%Y-%m-%d").date().isoformat()
    except ValueError as error:
        raise ImportValidationError(f"{label}: tanggal harus menggunakan format YYYY-MM-DD.") from error


def _number(value, label, allow_zero=False):
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ImportValidationError(f"{label}: nilai harus berupa angka.") from error
    if number < 0 or (number == 0 and not allow_zero):
        raise ImportValidationError(f"{label}: nilai harus lebih besar dari nol.")
    return number


def _rows(workbook, sheet_name, required_headers):
    if sheet_name not in workbook.sheetnames:
        raise ImportValidationError(f"Sheet '{sheet_name}' tidak ditemukan.")
    sheet = workbook[sheet_name]
    header_row = None
    headers = []
    for candidate in range(1, 11):
        candidate_headers = [_text(cell.value).lower() for cell in sheet[candidate]]
        if all(header in candidate_headers for header in required_headers):
            header_row = candidate
            headers = candidate_headers
            break
    if header_row is None:
        header_row = 1
        headers = [_text(cell.value).lower() for cell in sheet[1]]
    missing = [header for header in required_headers if header not in headers]
    if missing:
        raise ImportValidationError(
            f"Sheet '{sheet_name}' tidak memiliki kolom: {', '.join(missing)}."
        )
    indexes = {header: headers.index(header) for header in required_headers}
    result = []
    for row_number, values in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        if not any(value not in (None, "") for value in values):
            continue
        result.append((row_number, {header: values[indexes[header]] for header in required_headers}))
    return result


def import_workbook(file_storage):
    if not file_storage or not file_storage.filename:
        raise ImportValidationError("Pilih file Excel terlebih dahulu.")
    if not file_storage.filename.lower().endswith(".xlsx"):
        raise ImportValidationError("Format file harus .xlsx.")
    try:
        workbook = load_workbook(BytesIO(file_storage.read()), read_only=True, data_only=True)
    except Exception as error:
        raise ImportValidationError("File Excel tidak dapat dibaca atau rusak.") from error

    po_rows = _rows(workbook, "PurchaseOrders", SHEETS["PurchaseOrders"])
    receipt_rows = _rows(workbook, "GoodsReceipts", SHEETS["GoodsReceipts"])
    invoice_rows = _rows(workbook, "Invoices", SHEETS["Invoices"])
    if not po_rows and not receipt_rows and not invoice_rows:
        raise ImportValidationError("Workbook tidak memiliki baris data untuk diimpor.")

    db = get_db()
    counts = {"purchase_orders": 0, "goods_receipts": 0, "invoices": 0, "lines": 0}
    try:
        with db:
            po_cache = {}
            receipt_cache = {}
            invoice_cache = {}

            for row_number, row in po_rows:
                prefix = f"PurchaseOrders baris {row_number}"
                required = ["po_number", "supplier_code", "supplier_name", "product_sku", "product_name"]
                if any(not _text(row[field]) for field in required):
                    raise ImportValidationError(f"{prefix}: identitas PO, supplier, dan produk wajib diisi.")
                po_number = _text(row["po_number"])
                supplier_code = _text(row["supplier_code"])
                sku = _text(row["product_sku"])
                db.execute(
                    """INSERT INTO suppliers (code, name) VALUES (?, ?)
                       ON CONFLICT(code) DO UPDATE SET name=excluded.name""",
                    (supplier_code, _text(row["supplier_name"])),
                )
                supplier_id = db.execute("SELECT id FROM suppliers WHERE code=?", (supplier_code,)).fetchone()[0]
                db.execute(
                    """INSERT INTO products (sku, name) VALUES (?, ?)
                       ON CONFLICT(sku) DO UPDATE SET name=excluded.name""",
                    (sku, _text(row["product_name"])),
                )
                product_id = db.execute("SELECT id FROM products WHERE sku=?", (sku,)).fetchone()[0]

                if po_number not in po_cache:
                    if db.execute("SELECT 1 FROM purchase_orders WHERE po_number=?", (po_number,)).fetchone():
                        raise ImportValidationError(f"{prefix}: PO {po_number} sudah tersedia.")
                    cursor = db.execute(
                        """INSERT INTO purchase_orders
                           (po_number, supplier_id, order_date, expected_date, status)
                           VALUES (?, ?, ?, ?, 'Submitted')""",
                        (po_number, supplier_id, _date(row["order_date"], prefix),
                         _date(row["expected_date"], prefix) if row["expected_date"] else None),
                    )
                    po_cache[po_number] = (cursor.lastrowid, supplier_id)
                    db.execute(
                        """INSERT INTO transaction_history (entity_type, entity_id, action, details)
                           VALUES ('po', ?, 'IMPORTED', 'Dokumen dibuat melalui impor Excel sebagai Submitted.')""",
                        (cursor.lastrowid,),
                    )
                    counts["purchase_orders"] += 1
                po_id, expected_supplier_id = po_cache[po_number]
                if supplier_id != expected_supplier_id:
                    raise ImportValidationError(f"{prefix}: satu PO tidak boleh memiliki supplier berbeda.")
                db.execute(
                    "INSERT INTO po_items (purchase_order_id, product_id, quantity, unit_price) VALUES (?, ?, ?, ?)",
                    (po_id, product_id, _number(row["quantity"], prefix),
                     _number(row["unit_price"], prefix, allow_zero=True)),
                )
                counts["lines"] += 1

            for row_number, row in receipt_rows:
                prefix = f"GoodsReceipts baris {row_number}"
                receipt_number, po_number, sku = map(_text, (row["receipt_number"], row["po_number"], row["product_sku"]))
                if not all((receipt_number, po_number, sku)):
                    raise ImportValidationError(f"{prefix}: receipt_number, po_number, dan product_sku wajib diisi.")
                po = db.execute(
                    "SELECT id FROM purchase_orders WHERE po_number=? AND status='Submitted'", (po_number,)
                ).fetchone()
                product = db.execute("SELECT id FROM products WHERE sku=?", (sku,)).fetchone()
                if not po or not product:
                    raise ImportValidationError(f"{prefix}: PO atau SKU tidak ditemukan.")
                if receipt_number not in receipt_cache:
                    if db.execute("SELECT 1 FROM goods_receipts WHERE receipt_number=?", (receipt_number,)).fetchone():
                        raise ImportValidationError(f"{prefix}: goods receipt {receipt_number} sudah tersedia.")
                    cursor = db.execute(
                        "INSERT INTO goods_receipts (receipt_number, purchase_order_id, receipt_date, status) VALUES (?, ?, ?, 'Submitted')",
                        (receipt_number, po[0], _date(row["receipt_date"], prefix)),
                    )
                    receipt_cache[receipt_number] = (cursor.lastrowid, po[0])
                    db.execute(
                        """INSERT INTO transaction_history (entity_type, entity_id, action, details)
                           VALUES ('receipt', ?, 'IMPORTED', 'Dokumen dibuat melalui impor Excel sebagai Submitted.')""",
                        (cursor.lastrowid,),
                    )
                    counts["goods_receipts"] += 1
                receipt_id, expected_po_id = receipt_cache[receipt_number]
                if po[0] != expected_po_id:
                    raise ImportValidationError(f"{prefix}: satu goods receipt tidak boleh merujuk PO berbeda.")
                db.execute(
                    "INSERT INTO receipt_items (goods_receipt_id, product_id, quantity) VALUES (?, ?, ?)",
                    (receipt_id, product[0], _number(row["quantity"], prefix)),
                )
                counts["lines"] += 1

            for row_number, row in invoice_rows:
                prefix = f"Invoices baris {row_number}"
                invoice_number, po_number, supplier_code, sku = map(
                    _text, (row["invoice_number"], row["po_number"], row["supplier_code"], row["product_sku"])
                )
                if not all((invoice_number, po_number, supplier_code, sku)):
                    raise ImportValidationError(f"{prefix}: identitas invoice, PO, supplier, dan produk wajib diisi.")
                po = db.execute(
                    "SELECT id, supplier_id FROM purchase_orders WHERE po_number=? AND status='Submitted'",
                    (po_number,),
                ).fetchone()
                supplier = db.execute("SELECT id FROM suppliers WHERE code=?", (supplier_code,)).fetchone()
                product = db.execute("SELECT id FROM products WHERE sku=?", (sku,)).fetchone()
                if not po or not supplier or not product:
                    raise ImportValidationError(f"{prefix}: PO, supplier, atau SKU tidak ditemukan.")
                if invoice_number not in invoice_cache:
                    if db.execute("SELECT 1 FROM invoices WHERE invoice_number=?", (invoice_number,)).fetchone():
                        raise ImportValidationError(f"{prefix}: invoice {invoice_number} sudah tersedia.")
                    cursor = db.execute(
                        """INSERT INTO invoices
                           (invoice_number, purchase_order_id, supplier_id, invoice_date, status)
                           VALUES (?, ?, ?, ?, 'Submitted')""",
                        (invoice_number, po[0], supplier[0], _date(row["invoice_date"], prefix)),
                    )
                    invoice_cache[invoice_number] = (cursor.lastrowid, po[0], supplier[0])
                    db.execute(
                        """INSERT INTO transaction_history (entity_type, entity_id, action, details)
                           VALUES ('invoice', ?, 'IMPORTED', 'Dokumen dibuat melalui impor Excel sebagai Submitted.')""",
                        (cursor.lastrowid,),
                    )
                    counts["invoices"] += 1
                invoice_id, expected_po_id, expected_supplier_id = invoice_cache[invoice_number]
                if po[0] != expected_po_id or supplier[0] != expected_supplier_id:
                    raise ImportValidationError(f"{prefix}: satu invoice tidak boleh memiliki PO atau supplier berbeda.")
                db.execute(
                    "INSERT INTO invoice_items (invoice_id, product_id, quantity, unit_price) VALUES (?, ?, ?, ?)",
                    (invoice_id, product[0], _number(row["quantity"], prefix),
                     _number(row["unit_price"], prefix, allow_zero=True)),
                )
                counts["lines"] += 1
    except ImportValidationError:
        raise
    except Exception as error:
        raise ImportValidationError(f"Impor dibatalkan: {error}") from error
    finally:
        workbook.close()
    return counts
